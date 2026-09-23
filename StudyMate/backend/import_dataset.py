import os
import sys
import datetime
from pathlib import Path
from pymongo import MongoClient
from pypdf import PdfReader
import docx
from pptx import Presentation

# MongoDB Connection
client = MongoClient("mongodb://localhost:27017/")
db = client["studymate_rcpit"]
notes_col = db["notes"]
syllabus_col = db["syllabus"]

def extract_from_pdf(file_path):
    text = ""
    try:
        reader = PdfReader(file_path)
        for page in reader.pages:
            t = page.extract_text()
            if t:
                text += t + "\n"
    except Exception as e:
        print(f"  [!] PDF read issue in {file_path.name}: {e}")
    return text.strip()

def extract_from_docx(file_path):
    text = ""
    try:
        doc = docx.Document(file_path)
        for p in doc.paragraphs:
            if p.text:
                text += p.text + "\n"
    except Exception as e:
        print(f"  [!] DOCX read issue in {file_path.name}: {e}")
    return text.strip()

def extract_from_pptx(file_path):
    text = ""
    try:
        prs = Presentation(file_path)
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text:
                    text += shape.text + "\n"
    except Exception as e:
        print(f"  [!] PPTX read issue in {file_path.name}: {e}")
    return text.strip()

def extract_content(file_path):
    ext = file_path.suffix.lower()
    if ext == ".pdf":
        return extract_from_pdf(file_path)
    elif ext in [".docx", ".doc"]:
        return extract_from_docx(file_path)
    elif ext in [".pptx", ".ppt"]:
        return extract_from_pptx(file_path)
    elif ext in [".txt", ".md"]:
        try:
            return file_path.read_text(encoding="utf-8", errors="ignore")
        except Exception:
            return ""
    return ""

def guess_metadata(file_path, folder_path):
    """Detects academic metadata from file name or parent directories."""
    full_str = (str(file_path) + " " + file_path.stem).lower()
    
    # Branch
    branch = "AIML"
    if "cse" in full_str or "computer" in full_str:
        branch = "CSE"
    elif "ds" in full_str or "data science" in full_str:
        branch = "DS"
    elif "it" in full_str or "information" in full_str:
        branch = "IT"

    # Semester
    semester = "Semester 8"
    for s in range(1, 9):
        if f"sem {s}" in full_str or f"sem-{s}" in full_str or f"semester {s}" in full_str or f"sem{s}" in full_str:
            semester = f"Semester {s}"
            break

    # Unit
    unit = "Unit 1"
    for u in range(1, 7):
        if f"unit {u}" in full_str or f"unit-{u}" in full_str or f"unit{u}" in full_str or f"module {u}" in full_str:
            unit = f"Unit {u}"
            break

    # Type detection: Syllabus or Notes
    is_syllabus = "syllabus" in full_str or "curriculum" in full_str

    # Subject detection
    subject = "Deep Learning"
    if "machine learning" in full_str or " ml " in f" {full_str} ":
        subject = "Machine Learning"
    elif "dbms" in full_str or "database" in full_str:
        subject = "Database Systems"
    elif "network" in full_str or " cn " in f" {full_str} ":
        subject = "Computer Networks"
    elif "nlp" in full_str or "natural language" in full_str:
        subject = "Natural Language Processing"
    elif "cloud" in full_str:
        subject = "Cloud Computing"
    else:
        subject = file_path.parent.name if file_path.parent.name != folder_path.name else file_path.stem.replace("_", " ").title()

    return branch, semester, unit, subject, is_syllabus

def import_folder(target_folder_path):
    folder = Path(target_folder_path)
    if not folder.exists():
        print(f"\n[ERROR] Folder does not exist: {target_folder_path}")
        return

    supported_extensions = {".pdf", ".docx", ".doc", ".pptx", ".ppt", ".txt"}
    files_to_process = [p for p in folder.rglob("*") if p.suffix.lower() in supported_extensions]

    if not files_to_process:
        print(f"\n[!] No PDF, DOCX, or PPTX files found inside: {target_folder_path}")
        return

    print("\n=======================================================")
    print(f"Found {len(files_to_process)} document(s) in {folder.name}")
    print("Extracting text and indexing into MongoDB...")
    print("=======================================================\n")

    imported_count = 0
    for idx, fpath in enumerate(files_to_process, 1):
        content = extract_content(fpath)
        if not content:
            print(f"[{idx}/{len(files_to_process)}] ⚠️ Skipped (no text readable): {fpath.name}")
            continue

        branch, semester, unit, subject, is_syllabus = guess_metadata(fpath, folder)

        doc = {
            "title": fpath.stem.replace("_", " ").title(),
            "file_name": fpath.name,
            "branch": branch,
            "semester": semester,
            "subject": subject,
            "unit": unit,
            "content": content,
            "char_count": len(content),
            "file_type": fpath.suffix.lower().replace(".", "").upper(),
            "created_at": datetime.datetime.utcnow()
        }

        target_collection = syllabus_col if is_syllabus else notes_col
        target_collection.update_one({"title": doc["title"]}, {"$set": doc}, upsert=True)

        category_label = "Syllabus" if is_syllabus else "Notes"
        print(f"[{idx}/{len(files_to_process)}] ✅ Indexed to {category_label}: {doc['title']} ({subject} - {unit}) [{len(content)} chars]")
        imported_count += 1

    print(f"\n🎉 [COMPLETE] Successfully indexed {imported_count} documents into StudyMate MongoDB.")

if __name__ == "__main__":
    folder_input = input("\nEnter the full path to your dataset folder: ").strip().strip('"').strip("'")
    import_folder(folder_input)